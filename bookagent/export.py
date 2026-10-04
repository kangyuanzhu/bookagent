"""Readable, source-linked exports of a video book.

Each input page becomes exactly one slide, one PDF page, and one HTML article.
The HTML and PowerPoint notes always retain the complete generated text and
source transcript, even when a print layout has to shorten unusually long text.
"""

from __future__ import annotations

import base64
from dataclasses import asdict, is_dataclass
from functools import lru_cache
from html import escape
import json
import math
from pathlib import Path
from typing import TYPE_CHECKING, Any
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit

if TYPE_CHECKING:
    from .models import Book, Page

IVORY = "F7F4ED"
NAVY = "162B3D"
TEAL = "167C80"
MUTED = "637381"
BORDER = "DDDCD5"
WHITE = "FFFFFF"
WIDTH = 13.333333
HEIGHT = 7.5
IMAGE_BOX = (0.48, 1.58, 8.10, 4.95)
TEXT_BOX = (8.97, 1.61, 3.84, 4.98)
FONT_PATHS = (
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
)


def timestamp(seconds: float) -> str:
    """Format seconds as a stable, readable playback timestamp."""
    value = max(0, int(seconds))
    hours, rest = divmod(value, 3600)
    minutes, seconds = divmod(rest, 60)
    return f"{hours}:{minutes:02}:{seconds:02}" if hours else f"{minutes}:{seconds:02}"


def source_link(source: str, seconds: float, *, allow_local: bool = False) -> str | None:
    """Build a playback URL; discard YouTube playlist and tracking parameters.

    Arbitrary URL schemes never become browser links. Local files may become
    file: links only in the desktop PowerPoint/PDF exports.
    """
    if not source:
        return None
    try:
        parsed = urlsplit(source)
        host = parsed.hostname
    except ValueError:
        return None
    if parsed.scheme.lower() in {"http", "https"} and host:
        host = host.lower()
        video_id = None
        if host == "youtu.be":
            video_id = parsed.path.strip("/").split("/")[0]
        elif host == "youtube.com" or host.endswith(".youtube.com"):
            if parsed.path.rstrip("/") == "/watch":
                video_id = parse_qs(parsed.query).get("v", [None])[0]
            elif parsed.path.startswith(("/shorts/", "/embed/", "/live/")):
                video_id = parsed.path.split("/")[2]
        value = str(max(0, int(seconds)))
        if video_id:
            return "https://www.youtube.com/watch?" + urlencode({"v": video_id, "t": value})
        query = parse_qs(parsed.query, keep_blank_values=True)
        query["t"] = [value]
        return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urlencode(query, doseq=True), ""))
    if allow_local and not parsed.scheme:
        path = Path(source).expanduser()
        if path.is_file():
            return path.resolve().as_uri() + f"#t={max(0, int(seconds))}"
    return None


def _warn(book: Any, warning: str) -> None:
    warnings = getattr(book, "warnings", None)
    if warnings is not None and warning not in warnings:
        warnings.append(warning)


def _image_path(page: Any) -> Path | None:
    source = getattr(page, "frame_path", None)
    if source:
        path = Path(source)
        if path.is_file():
            return path
    return None


def _sections(page: Any, *, overview: bool = False) -> list[tuple[str, str]]:
    sections = []
    if page.summary:
        kind = getattr(page, "kind", "content")
        label = {"cover": "READING TARGET", "overview": "AT A GLANCE"}.get(kind, "PRESENTER SUMMARY")
        if kind == "content" and "no ai summary" in getattr(page, "uncertainty", "").casefold():
            label = "TRANSCRIPT EXCERPT"
        sections.append((label, page.summary))
    if page.bullets and not overview:
        sections.append(("KEY IDEAS", "\n".join(f"• {bullet}" for bullet in page.bullets)))
    if page.takeaway:
        sections.append(("TAKEAWAY", page.takeaway))
    if getattr(page, "uncertainty", ""):
        sections.append(("SOURCE NOTE", page.uncertainty))
    return sections


def _notes(book: Any, page: Any) -> str:
    source = source_link(book.source_url, page.start, allow_local=True) or book.source_url
    parts = [page.title, f"Source interval: {timestamp(page.start)} – {timestamp(page.end)}", f"Source: {source}"]
    for heading, body in _sections(page):
        parts.extend(["", heading, body])
    source_heading = "SOURCE REFERENCE" if getattr(book, "metadata", {}).get("embed_transcript") is False else "FULL SOURCE TRANSCRIPT"
    parts.extend(["", source_heading, page.source_text or "No source transcript was provided for this page."])
    return "\n".join(parts)


@lru_cache(maxsize=1)
def _font_path() -> str | None:
    return next((path for path in FONT_PATHS if Path(path).is_file()), None)


@lru_cache(maxsize=64)
def _measurement_font(size: float) -> tuple[Any, int]:
    from PIL import ImageFont

    font_path = _font_path()
    if font_path:
        return ImageFont.truetype(font_path, max(1, int(size * 4))), 4
    return ImageFont.load_default(), 1


def _wrapped_lines(text: str, width: float, size: float) -> int:
    """Estimate actual font widths in points, including unbroken long strings."""
    font, scale = _measurement_font(size)
    max_width = width * 72 * scale
    lines = 0
    for paragraph in text.split("\n"):
        if not paragraph:
            lines += 1
            continue
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if current and font.getlength(candidate) > max_width:
                lines += 1
                current = ""
            if font.getlength(word) > max_width:
                if current:
                    lines += 1
                    current = ""
                chunk = ""
                for character in word:
                    if chunk and font.getlength(chunk + character) > max_width:
                        lines += 1
                        chunk = character
                    else:
                        chunk += character
                current = chunk
            else:
                current = f"{current} {word}".strip()
        lines += bool(current)
    return max(1, lines)


def _shorten(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    clipped = text[: max(1, limit - 1)]
    boundary = clipped.rfind(" ")
    if boundary > len(clipped) * 0.6:
        clipped = clipped[:boundary]
    return clipped.rstrip(" .\n") + "…"


def _fit_title(title: str, width: float = 12.3, height: float = 0.76, size: float = 27) -> tuple[str, float, bool]:
    original = title
    # Bound measurement work for malformed or unusually verbose model output.
    title = _shorten(title, 1000)
    while size > 19 and _wrapped_lines(title, width, size) * size * 1.14 > height * 72:
        size -= 1
    while _wrapped_lines(title, width, size) * size * 1.14 > height * 72 and len(title) > 20:
        title = _shorten(original, int(len(title) * 0.9))
    return title, size, title != original


def _fit_sections(sections: list[tuple[str, str]], width: float = 3.84, height: float = 4.98) -> tuple[list[tuple[str, str]], float, bool]:
    original = sections
    sections = [(heading, _shorten(body, 1600)) for heading, body in sections]
    size = 16.0

    def needed(values: list[tuple[str, str]], font_size: float) -> float:
        return sum(14 + 9 + _wrapped_lines(body, width, font_size) * font_size * 1.18 + body.count("\n") * 3 + 16 for _, body in values)

    while size > 12 and needed(sections, size) > height * 72:
        size -= 0.5
    limit = sum(len(body) for _, body in sections)
    while needed(sections, size) > height * 72 and limit > 80:
        limit = int(limit * 0.87)
        total = max(1, sum(len(body) for _, body in original))
        sections = [(head, _shorten(body, max(28, int(limit * len(body) / total)))) for head, body in original]
    return sections, size, sections != original


def _footnote(page: Any, shortened: bool) -> str:
    if shortened:
        return "Full text is available in HTML and PowerPoint notes."
    if getattr(page, "uncertainty", ""):
        return "Check the source note before relying on this summary."
    return "Read the summary. Return to the source for detail."


def _export_pptx(book: Any, path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    deck.core_properties.title = book.title
    deck.core_properties.subject = "Video book with presenter summaries and source timestamps"
    deck.core_properties.author = "Bookagent"

    def rectangle(slide: Any, box: tuple[float, float, float, float], color: str, border: str | None = None) -> Any:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(value) for value in box))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        if border:
            shape.line.color.rgb = RGBColor.from_string(border)
        else:
            shape.line.fill.background()
        return shape

    def textbox(slide: Any, text: str, box: tuple[float, float, float, float], size: float, color: str = NAVY, bold: bool = False, link: str | None = None) -> Any:
        shape = slide.shapes.add_textbox(*(Inches(value) for value in box))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.line_spacing = 1.14
            paragraph.space_after = Pt(0)
            if link:
                for run in paragraph.runs:
                    run.hyperlink.address = link
        return shape

    for page in book.pages:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        rectangle(slide, (0, 0, WIDTH, HEIGHT), IVORY)
        rectangle(slide, (0, 0, WIDTH, 0.09), TEAL)
        chapter = getattr(page, "chapter", "")
        textbox(slide, f"VIDEO BOOK   /   {chapter or book.title}"[:140], (0.48, 0.34, 12.1, 0.2), 9, MUTED)
        title, title_size, shortened = _fit_title(page.title)
        textbox(slide, title, (0.48, 0.69, 12.33, 0.76), title_size, bold=True)
        rectangle(slide, IMAGE_BOX, WHITE, BORDER)
        kind = getattr(page, "kind", "content")
        image = _image_path(page)
        if image:
            with Image.open(image) as picture:
                iw, ih = picture.size
            x, y, w, h = IMAGE_BOX
            ratio = min((w - 0.16) / iw, (h - 0.16) / ih)
            pw, ph = iw * ratio, ih * ratio
            slide.shapes.add_picture(str(image), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph))
        elif kind == "cover":
            rectangle(slide, IMAGE_BOX, NAVY)
            cover_title, cover_size, cover_short = _fit_title(book.title, 6.7, 2.2, 35)
            shortened |= cover_short
            textbox(slide, "A VIDEO, DISTILLED", (1.05, 2.15, 6.5, 0.3), 12, "70CDD0", True)
            textbox(slide, cover_title, (1.05, 2.72, 6.7, 2.2), cover_size, WHITE, True)
            textbox(slide, f"{timestamp(book.duration)} of video  ·  {len(book.pages)} pages\nReading target: {book.reading_minutes:g} minutes", (1.05, 5.55, 6.7, 0.62), 16, "D6E2E7")
        elif kind == "overview":
            textbox(slide, "YOUR READING MAP", (0.92, 2.01, 7.2, 0.35), 12, TEAL, True)
            overview_text = "\n\n".join(page.bullets) or page.summary
            values, size, overview_short = _fit_sections([("", overview_text)], 7.15, 3.75)
            shortened |= overview_short
            textbox(slide, values[0][1], (0.92, 2.55, 7.15, 3.75), size)
        else:
            textbox(slide, "Transcript-only page", (1.1, 3.33, 6.84, 0.6), 26, TEAL, True)
            textbox(slide, "No source slide was captured for this interval.", (1.1, 4.06, 6.84, 0.7), 16, MUTED)
        sections, size, body_short = _fit_sections(_sections(page, overview=kind == "overview"))
        shortened |= body_short
        shape = slide.shapes.add_textbox(*(Inches(value) for value in TEXT_BOX))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        first = True
        for heading, body in sections:
            label = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            label.text = heading
            label.font.name, label.font.size = "Arial", Pt(10)
            label.font.bold = True
            label.font.color.rgb = RGBColor.from_string(TEAL)
            label.space_after = Pt(7)
            for line in body.split("\n"):
                paragraph = frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name, paragraph.font.size = "Arial", Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(NAVY)
                paragraph.line_spacing = 1.18
                paragraph.space_after = Pt(3)
            paragraph.space_after = Pt(16)
        link = source_link(book.source_url, page.start, allow_local=True)
        textbox(slide, f"{timestamp(page.start)} – {timestamp(page.end)}  ↗ source", (0.48, 6.99, 4.2, 0.23), 10, TEAL, link=link)
        textbox(slide, _footnote(page, shortened), (4.5, 6.99, 6.9, 0.23), 9, MUTED)
        textbox(slide, f"{page.number:03} / {len(book.pages):03}", (11.78, 6.98, 1.1, 0.25), 10, NAVY, True)
        slide.notes_slide.notes_text_frame.text = _notes(book, page)
        if shortened:
            _warn(book, f"Page {page.number}: some print-layout text was shortened; full text remains in HTML and PowerPoint notes.")
    deck.save(path)


def _export_pdf(book: Any, path: Path) -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import Paragraph

    font_name = "Helvetica"
    path_to_font = _font_path()
    if path_to_font:
        font_name = "BookagentUnicode"
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, path_to_font))
        supported = pdfmetrics.getFont(font_name).face.charToGlyph
    else:
        supported = None
        _warn(book, "No Unicode PDF font was available. Non-Latin text may require a system font; HTML and PowerPoint retain the original text.")

    def safe(text: str) -> str:
        unsupported = {character for character in text if not character.isspace() and ((supported is not None and ord(character) not in supported) or (supported is None and ord(character) > 255))}
        if unsupported:
            _warn(book, "The PDF font does not contain every source character. Unsupported glyphs are shown as '?' in PDF; HTML and PowerPoint retain the originals.")
            return "".join("?" if character in unsupported else character for character in text)
        return text

    canvas = Canvas(str(path), pagesize=(WIDTH * 72, HEIGHT * 72), pageCompression=1)
    canvas.setTitle(book.title)
    canvas.setAuthor("Bookagent")

    def rect(box: tuple[float, float, float, float], color: str, border: str | None = None) -> None:
        x, y, w, h = box
        canvas.setFillColor(HexColor("#" + color))
        if border:
            canvas.setStrokeColor(HexColor("#" + border))
        canvas.rect(x * 72, (HEIGHT - y - h) * 72, w * 72, h * 72, fill=1, stroke=bool(border))

    def paragraph(text: str, box: tuple[float, float, float, float], size: float, color: str = NAVY, *, fit: bool = False) -> bool:
        x, y, w, h = box
        original = text
        while True:
            style = ParagraphStyle("book", fontName=font_name, fontSize=size, leading=size * 1.18, textColor=HexColor("#" + color), splitLongWords=True)
            flow = Paragraph(escape(safe(text)).replace("\n", "<br/>"), style)
            _, used = flow.wrap(w * 72, h * 72)
            if used <= h * 72:
                flow.drawOn(canvas, x * 72, (HEIGHT - y) * 72 - used)
                return text != original
            if size > 11 and fit:
                size -= 0.5
            elif fit and len(text) > 8:
                text = _shorten(original, max(8, int(len(text) * 0.9)))
            else:
                raise ValueError(f"PDF text exceeds its layout box: {text[:80]!r}")

    for page in book.pages:
        rect((0, 0, WIDTH, HEIGHT), IVORY)
        rect((0, 0, WIDTH, 0.09), TEAL)
        paragraph(f"VIDEO BOOK / {getattr(page, 'chapter', '') or book.title}"[:140], (0.48, 0.34, 12.1, 0.25), 9, MUTED, fit=True)
        title, size, shortened = _fit_title(page.title)
        shortened |= paragraph(title, (0.48, 0.69, 12.33, 0.76), size, fit=True)
        rect(IMAGE_BOX, WHITE, BORDER)
        image = _image_path(page)
        kind = getattr(page, "kind", "content")
        if image:
            x, y, w, h = IMAGE_BOX
            canvas.drawImage(ImageReader(str(image)), (x + 0.08) * 72, (HEIGHT - y - h + 0.08) * 72, (w - 0.16) * 72, (h - 0.16) * 72, preserveAspectRatio=True, anchor="c", mask="auto")
        elif kind == "cover":
            rect(IMAGE_BOX, NAVY)
            paragraph("A VIDEO, DISTILLED", (1.05, 2.15, 6.5, 0.35), 12, "70CDD0")
            shortened |= paragraph(book.title, (1.05, 2.72, 6.7, 2.2), 35, WHITE, fit=True)
            paragraph(f"{timestamp(book.duration)} of video · {len(book.pages)} pages\nReading target: {book.reading_minutes:g} minutes", (1.05, 5.55, 6.7, 0.65), 16, "D6E2E7", fit=True)
        elif kind == "overview":
            paragraph("YOUR READING MAP", (0.92, 2.01, 7.2, 0.35), 12, TEAL)
            shortened |= paragraph("\n\n".join(page.bullets) or page.summary, (0.92, 2.55, 7.15, 3.75), 19, fit=True)
        else:
            paragraph("Transcript-only page", (1.1, 3.33, 6.84, 0.6), 26, TEAL)
            paragraph("No source slide was captured for this interval.", (1.1, 4.06, 6.84, 0.7), 16, MUTED)
        sections, size, body_short = _fit_sections(_sections(page, overview=kind == "overview"))
        shortened |= body_short
        # Fit the complete group using the PDF engine, whose line breaking can
        # differ from PowerPoint. Never let a block overwrite the footer.
        def render_blocks(values: list[tuple[str, str]], body_size: float, draw: bool = False) -> float:
            y = TEXT_BOX[1] * 72
            for heading, body in values:
                for text, fs, color, after in ((heading, 10, TEAL, 7), (body, body_size, NAVY, 16)):
                    style = ParagraphStyle("section", fontName=font_name, fontSize=fs, leading=fs * 1.18, textColor=HexColor("#" + color), splitLongWords=True)
                    flow = Paragraph(escape(safe(text)).replace("\n", "<br/>"), style)
                    _, used = flow.wrap(TEXT_BOX[2] * 72, TEXT_BOX[3] * 72)
                    if draw:
                        flow.drawOn(canvas, TEXT_BOX[0] * 72, HEIGHT * 72 - y - used)
                    y += used + after
            return y - TEXT_BOX[1] * 72

        original_sections = sections
        total_chars = sum(len(body) for _, body in sections)
        while render_blocks(sections, size) > TEXT_BOX[3] * 72:
            if size > 11:
                size -= 0.5
            else:
                total_chars = int(total_chars * 0.85)
                if total_chars < 20:
                    raise ValueError(f"Page {page.number} contains too many summary sections to fit.")
                total = max(1, sum(len(body) for _, body in original_sections))
                sections = [(heading, _shorten(body, max(8, int(total_chars * len(body) / total)))) for heading, body in original_sections]
                shortened = True
        render_blocks(sections, size, draw=True)
        paragraph(f"{timestamp(page.start)} – {timestamp(page.end)} / source", (0.48, 6.99, 4.0, 0.25), 10, TEAL)
        link = source_link(book.source_url, page.start, allow_local=True)
        if link:
            canvas.linkURL(link, (0.48 * 72, 0.22 * 72, 4.3 * 72, 0.51 * 72), relative=0, thickness=0)
        paragraph(_footnote(page, shortened), (4.5, 6.99, 6.9, 0.25), 9, MUTED, fit=True)
        paragraph(f"{page.number:03} / {len(book.pages):03}", (11.78, 6.98, 1.1, 0.25), 10)
        if shortened:
            _warn(book, f"Page {page.number}: some print-layout text was shortened; full text remains in HTML and PowerPoint notes.")
        canvas.showPage()
    canvas.save()


def _data_image(path: Path) -> str:
    from PIL import Image
    import io

    data = path.read_bytes()
    with Image.open(path) as picture:
        image_format = picture.format
        if image_format not in {"PNG", "JPEG", "GIF", "WEBP"}:
            output = io.BytesIO()
            picture.convert("RGB").save(output, format="PNG")
            data, image_format = output.getvalue(), "PNG"
    mime = {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif", "WEBP": "image/webp"}[image_format]
    return f"data:{mime};base64," + base64.b64encode(data).decode("ascii")


def _export_html(book: Any, path: Path) -> None:
    articles = []
    for page in book.pages:
        image = _image_path(page)
        kind = getattr(page, "kind", "content")
        if image:
            visual = f'<img src="{_data_image(image)}" alt="Captured source visual for interval {timestamp(page.start)}–{timestamp(page.end)}" loading="lazy">'
        elif kind == "cover":
            visual = f'<div class="cover"><span>A VIDEO, DISTILLED</span><h2>{escape(book.title)}</h2><p>{timestamp(book.duration)} of video · {len(book.pages)} pages<br>Reading target: {book.reading_minutes:g} minutes</p></div>'
        elif kind == "overview":
            visual = '<div class="overview"><span>YOUR READING MAP</span><ol>' + "".join(f"<li>{escape(bullet)}</li>" for bullet in page.bullets) + "</ol></div>"
        else:
            visual = '<div class="missing"><strong>Transcript-only page</strong><p>No source slide was captured for this interval.</p></div>'
        sections = "".join(f'<section class="summary-section"><h3>{escape(heading)}</h3><p>{escape(body).replace(chr(10), "<br>")}</p></section>' for heading, body in _sections(page, overview=kind == "overview"))
        url = source_link(book.source_url, page.start)
        label = f"{timestamp(page.start)} – {timestamp(page.end)}"
        playback = f'<a href="{escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">{label} ↗ source</a>' if url else f'<span>{label} · {escape(book.source_url or "Source unavailable")}</span>'
        source = escape(page.source_text or "No source transcript was provided for this page.")
        chapter = escape(getattr(page, "chapter", "") or book.title)
        source_label = "Source reference" if getattr(book, "metadata", {}).get("embed_transcript") is False else "Full source transcript"
        articles.append(f'<article class="book-page" id="page-{page.number}" tabindex="-1"><div class="eyebrow">VIDEO BOOK / {chapter}</div><h2 class="page-title">{escape(page.title)}</h2><div class="page-grid"><div class="visual">{visual}</div><aside>{sections}</aside></div><footer>{playback}<span>{page.number:03} / {len(book.pages):03}</span></footer><details><summary>{source_label} · {label}</summary><pre>{source}</pre></details></article>')
    warning_html = "".join(f"<li>{escape(warning)}</li>" for warning in book.warnings)
    warnings = f'<details class="warnings"><summary>Source and export notes ({len(book.warnings)})</summary><ul>{warning_html}</ul></details>' if warning_html else ""
    document = '''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>''' + escape(book.title) + '''</title>
<style>
:root{--ink:#162b3d;--teal:#167c80;--paper:#f7f4ed;--muted:#637381;--line:#ddddd5}*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:90px}body{margin:0;background:#e7e9e8;color:var(--ink);font:16px/1.55 Arial,sans-serif}a{color:var(--teal)}header{position:sticky;top:0;z-index:5;display:flex;align-items:center;gap:20px;background:var(--ink);color:#fff;padding:15px max(20px,calc((100% - 1280px)/2));box-shadow:0 3px 15px #0002}header strong{font-size:17px;max-width:45%;line-height:1.25}input{padding:11px 14px;border:0;border-radius:6px;font:inherit;flex:1;min-width:100px}#counter{font-size:13px;white-space:nowrap}main{max-width:1320px;margin:28px auto;padding:0 20px}.book-page{background:var(--paper);border-top:6px solid var(--teal);padding:32px 40px 22px;margin-bottom:28px;box-shadow:0 8px 28px #162b3d14;scroll-margin-top:90px}.eyebrow{font-size:11px;letter-spacing:.14em;color:var(--muted);text-transform:uppercase}.page-title{font-size:28px;line-height:1.25;margin:16px 0 28px;overflow-wrap:anywhere}.page-grid{display:grid;grid-template-columns:minmax(0,2.1fr) minmax(230px,1fr);gap:36px}.visual{min-height:370px;display:flex;align-items:center;justify-content:center;background:#fff;border:1px solid var(--line)}.visual img{width:100%;max-height:540px;object-fit:contain}.summary-section{margin-bottom:23px}.summary-section h3{font-size:11px;letter-spacing:.12em;color:var(--teal);margin:0 0 9px}.summary-section p{margin:0;overflow-wrap:anywhere}.missing{text-align:center;padding:35px;color:var(--muted)}.missing strong{font-size:28px;color:var(--teal)}.cover{align-self:stretch;width:100%;background:var(--ink);color:#fff;padding:44px}.cover span,.overview>span{font-size:12px;letter-spacing:.14em;color:#70cdd0}.cover h2{font-size:34px;line-height:1.25;margin:28px 0 48px;overflow-wrap:anywhere}.cover p{color:#d6e2e7}.overview{padding:36px;align-self:flex-start;width:100%}.overview>span{color:var(--teal)}.overview li{padding:12px 0;font-size:20px}footer{display:flex;justify-content:space-between;gap:20px;border-top:1px solid var(--line);padding-top:16px;margin-top:27px;font-size:12px;color:var(--muted)}footer a{text-decoration:none}details{font-size:13px;margin-top:20px}summary{cursor:pointer;color:var(--teal)}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.65 Arial,sans-serif;background:#fff;padding:20px;border:1px solid var(--line)}.warnings{background:var(--paper);padding:18px 28px}.empty{padding:60px;text-align:center}button{background:var(--teal);border:0;color:#fff;font:inherit;border-radius:4px;padding:8px 14px;cursor:pointer}nav{display:flex;justify-content:center;gap:15px;margin:24px}small{color:var(--muted)}[hidden]{display:none!important}@media(max-width:800px){header{flex-wrap:wrap;gap:10px}header strong{max-width:100%;width:100%}.book-page{padding:24px 20px}.page-grid{grid-template-columns:1fr;gap:24px}.visual{min-height:250px}.page-title{font-size:24px}.cover h2{font-size:28px}.cover{padding:30px}.overview li{font-size:17px}}@media print{body{background:#fff}header,nav,.warnings,details{display:none}.book-page{break-after:page;box-shadow:none;margin:0;padding:20px}.page-grid{grid-template-columns:2fr 1fr}.visual{min-height:0}}
</style></head><body><header><strong>''' + escape(book.title) + '''</strong><input id="search" type="search" aria-label="Search summaries and transcripts" placeholder="Search summaries and transcripts…"><span id="counter"></span></header><main>''' + warnings + "".join(articles) + '''<p id="empty" class="empty" hidden>No pages match this search.</p><nav aria-label="Page navigation"><button id="previous">← Previous</button><button id="next">Next →</button></nav><p style="text-align:center"><small>Use ← and → to move between pages. Each source timestamp opens the original video.</small></p></main><script>
const pages=[...document.querySelectorAll('.book-page')],search=document.querySelector('#search'),counter=document.querySelector('#counter');let active=0;
function visible(){return pages.filter(p=>!p.hidden)}function update(){const query=search.value.toLocaleLowerCase().trim();pages.forEach(p=>p.hidden=query&&!p.textContent.toLocaleLowerCase().includes(query));const list=visible();counter.textContent=list.length+' / '+pages.length+' pages';document.querySelector('#empty').hidden=list.length>0;active=Math.min(active,Math.max(0,list.length-1))}function move(delta){const list=visible();if(!list.length)return;const closest=list.reduce((best,p,i)=>Math.abs(p.getBoundingClientRect().top-90)<best.distance?{index:i,distance:Math.abs(p.getBoundingClientRect().top-90)}:best,{index:active,distance:Infinity});active=Math.min(list.length-1,Math.max(0,closest.index+delta));list[active].scrollIntoView({behavior:'smooth',block:'start'});list[active].focus({preventScroll:true})}search.addEventListener('input',update);document.querySelector('#previous').onclick=()=>move(-1);document.querySelector('#next').onclick=()=>move(1);document.addEventListener('keydown',event=>{if(['INPUT','TEXTAREA'].includes(document.activeElement.tagName))return;if(event.key==='ArrowRight'){event.preventDefault();move(1)}if(event.key==='ArrowLeft'){event.preventDefault();move(-1)}});update();
</script></body></html>'''
    path.write_text(document, encoding="utf-8")


def _json_book(book: Any) -> dict[str, Any]:
    if is_dataclass(book):
        return asdict(book)
    # Also support compatible objects in integrations and lightweight fixtures.
    data = dict(vars(book))
    data["pages"] = [asdict(page) if is_dataclass(page) else dict(vars(page)) for page in book.pages]
    return data


def export_book(book: Book, outdir: Path, formats: tuple[str, ...] = ("pptx", "pdf", "html")) -> dict[str, Path]:
    """Write a book and its reusable JSON manifest to ``outdir``.

    No network calls are made. Fonts and images come from local files. At least
    one page is required, and unsupported formats fail before writing files.
    """
    unknown = set(formats) - {"pptx", "pdf", "html", "json"}
    if unknown:
        raise ValueError("Unsupported export format(s): " + ", ".join(sorted(unknown)))
    if not book.pages:
        raise ValueError("A book must contain at least one page.")
    transcript_only_count = 0
    for page in book.pages:
        if not all(math.isfinite(float(value)) and value >= 0 for value in (page.start, page.end)) or page.end < page.start:
            raise ValueError(f"Page {page.number} has an invalid source interval.")
        if getattr(page, "kind", "content") == "content" and not _image_path(page):
            if getattr(page, "frame_path", None):
                _warn(book, f"Page {page.number}: captured image file is missing; exported as a transcript-only page.")
            else:
                transcript_only_count += 1
    if transcript_only_count:
        _warn(book, f"{transcript_only_count} content page(s) have no captured source image and are exported as transcript-only pages.")
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    exports = {}
    for format_name, exporter in (("pptx", _export_pptx), ("pdf", _export_pdf), ("html", _export_html)):
        if format_name in formats:
            path = outdir / f"book.{format_name}"
            exporter(book, path)
            exports[format_name] = path
    manifest = outdir / "book.json"
    manifest.write_text(json.dumps(_json_book(book), ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    exports["json"] = manifest
    return exports
