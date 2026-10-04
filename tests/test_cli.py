import json
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from bookagent import cli, demo
from bookagent.models import timestamp
from bookagent.planner import count_words


def arguments(tmp_path, *options):
    return cli._parser().parse_args([
        "build", "--transcript", str(tmp_path / "missing.json"), "--provider", "extractive",
        "--out", str(tmp_path / "book"), *options,
    ])


@pytest.mark.parametrize("options, message", [
    (("--pages", "2"), "at least 3"),
    (("--reading-minutes", "0"), "reading-minutes"),
    (("--reading-minutes", "nan"), "reading-minutes"),
    (("--frame-interval", "0"), "frame-interval"),
    (("--workers", "9"), "workers"),
    (("--formats", "pptx,unknown"), "formats"),
    (("--duration", "-1"), "duration"),
    (("--crop", "0,0,0,1"), "crop"),
    (("--crop", "0.5,0,0.6,1"), "crop"),
    (("--crop", "0,0,nan,1"), "crop"),
    (("--crop", "0,0,1"), "crop"),
    (("--crop", "bad,0,1,1"), "crop"),
])
def test_invalid_options_fail_before_reading_input_or_creating_output(tmp_path, options, message):
    with pytest.raises(ValueError, match=message):
        cli.build_book(arguments(tmp_path, *options))
    assert not (tmp_path / "book").exists()


def test_missing_api_key_is_reported_before_input_access(tmp_path, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    with pytest.raises(ValueError, match="OPENAI_API_KEY"):
        cli.build_book(arguments(tmp_path, "--provider", "openai"))
    assert not (tmp_path / "book").exists()


def _write_transcript(path: Path, count: int = 8) -> list[dict]:
    segments = [
        {"start": index * 12, "end": index * 12 + 11.5,
         "text": " ".join(f"lesson{index}word{word}" for word in range(80))}
        for index in range(count)
    ]
    path.write_text(json.dumps(segments), encoding="utf-8")
    return segments


def test_extractive_full_pipeline_preserves_source_without_api_or_frames(tmp_path, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    transcript = tmp_path / "captions.json"
    segments = _write_transcript(transcript)
    out = tmp_path / "book"
    args = cli._parser().parse_args([
        "build", "--transcript", str(transcript), "--duration", "96", "--pages", "10",
        "--reading-minutes", "5", "--provider", "extractive", "--out", str(out),
    ])
    book = cli.build_book(args)
    assert len(book.pages) == 10
    content = book.pages[2:]
    assert [page.number for page in book.pages] == list(range(1, 11))
    assert content[0].start == 0 and content[-1].end == 96
    assert all(left.end == right.start for left, right in zip(content, content[1:]))
    assert "\n".join(page.source_text for page in content) == "\n".join(
        f"[{timestamp(segment['start'])}] {segment['text']}" for segment in segments
    )
    assert all(page.frame_path is None for page in content)
    assert all("Excerpt; no AI summary." == page.uncertainty for page in content)
    assert book.metadata["provider"] == "extractive"
    assert book.metadata["requested_pages"] == 10
    assert any("Transcript-only" in warning for warning in book.warnings)

    presentation = Presentation(out / "book.pptx")
    assert len(presentation.slides) == 10
    for page, slide in zip(content, list(presentation.slides)[2:]):
        assert page.source_text in slide.notes_slide.notes_text_frame.text
    reader = PdfReader(out / "book.pdf")
    assert len(reader.pages) == 10
    assert "Transcript" in reader.pages[2].extract_text()
    html = (out / "book.html").read_text(encoding="utf-8")
    assert "Transcript-only page" in html
    assert all(segment["text"] in html for segment in segments)
    manifest = json.loads((out / "book.json").read_text(encoding="utf-8"))
    assert len(manifest["pages"]) == 10
    assert manifest["metadata"]["source_words"] == 640


def test_reading_budget_error_prevents_incomplete_exports(tmp_path):
    transcript = tmp_path / "captions.json"
    _write_transcript(transcript)
    with pytest.raises(ValueError, match="Reading target is too short"):
        cli.build_book(arguments(tmp_path, "--transcript", str(transcript), "--pages", "10", "--reading-minutes", "0.1"))
    assert not (tmp_path / "book" / "book.json").exists()


def test_blank_transcript_reports_a_useful_error(tmp_path, capsys):
    transcript = tmp_path / "blank.json"
    transcript.write_text("[]", encoding="utf-8")
    status = cli.main(["build", "--transcript", str(transcript), "--provider", "extractive", "--out", str(tmp_path / "book")])
    assert status == 1
    assert "no usable timestamped captions" in capsys.readouterr().err
    assert not (tmp_path / "book" / "book.pptx").exists()


def test_demo_is_explicitly_synthetic_and_supports_missing_ffmpeg(tmp_path, monkeypatch):
    monkeypatch.setattr(demo.shutil, "which", lambda command: None)
    calls = []
    monkeypatch.setattr(cli, "build_book", lambda args: calls.append(args))
    out = tmp_path / "demo"
    demo.run_demo(out)
    source = out / "work" / "demo-source"
    assert len(list(source.glob("synthetic-slide-*.png"))) == 8
    saved = json.loads((source / "synthetic-transcript.json").read_text(encoding="utf-8"))
    assert saved["synthetic"] is True
    assert "not captions" in saved["description"]
    assert len(saved["segments"]) == 8
    assert all(count_words(segment["text"]) >= 75 for segment in saved["segments"])
    assert saved["segments"][-1]["end"] < 96
    assert len(calls) == 1
    args = calls[0]
    assert args.video is None
    assert args.provider == "extractive" and args.pages == 10
    assert args.reading_minutes == 5 and args.duration == 96
    assert args.out == out.resolve()
    assert "synthetic sample" in args.title


def test_concat_paths_quote_spaces_and_apostrophes(tmp_path):
    path = tmp_path / "reader's sample" / "slide 1.png"
    quoted = demo._concat_quote(path)
    assert quoted.startswith("'") and quoted.endswith("'")
    assert "reader'\\''s sample" in quoted
    assert "slide 1.png" in quoted
