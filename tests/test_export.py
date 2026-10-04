import json
from pathlib import Path

import pytest

from bookagent.export import export_book, source_link, timestamp
from bookagent.models import Book, Page


def sample_book(tmp_path: Path) -> Book:
    from PIL import Image

    frame = tmp_path / "source.png"
    Image.new("RGB", (1200, 675), "#167c80").save(frame)
    return Book(
        title="A lecture <with> examples & sources",
        source_url="https://www.youtube.com/watch?v=Tcn5Yb2K0h4&list=PLVs8SZOx0kX8&index=1",
        duration=14_400,
        pages=[
            Page(1, "A lecture, distilled", 0, 0, "Read the ideas with their source slides.", ["60 minutes", "Keep source context"], "Use timestamps to verify claims.", kind="cover"),
            Page(2, "The reading map", 0, 0, "Follow the sequence.", ["Introduction", "Worked examples", "Conclusions"], "Read the recap first.", kind="overview"),
            Page(3, "A captured source slide", 72, 120, "The presenter explains a specific example.", ["A practical idea", "A concrete limitation"], "Apply the idea after checking the source.", str(frame), source_text="<script>alert('never execute')</script>\n" + "complete source transcript " * 300, chapter="Examples"),
            Page(4, "A spoken explanation", 120, 180, "The presenter clarifies what the slide omits.", ["Keep assumptions explicit"], "Distinguish evidence from inference.", source_text="A transcript-only segment.", uncertainty="Captions may contain recognition errors."),
        ],
    )


def test_source_link_removes_playlist_and_sets_timestamp():
    assert source_link("https://www.youtube.com/watch?v=abc&list=123&index=3&t=10", 72.9) == "https://www.youtube.com/watch?v=abc&t=72"
    assert source_link("https://youtu.be/abc?si=tracking", 72) == "https://www.youtube.com/watch?v=abc&t=72"
    assert source_link("https://www.youtube.com/shorts/abc", 0) == "https://www.youtube.com/watch?v=abc&t=0"
    assert source_link("javascript:alert(1)", 0) is None
    assert source_link("data:text/html,hello", 0) is None
    assert timestamp(3723) == "1:02:03"


def test_exports_preserve_page_count_transcript_and_original_image(tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pypdf import PdfReader

    book = sample_book(tmp_path)
    exports = export_book(book, tmp_path / "output")
    assert set(exports) == {"pptx", "pdf", "html", "json"}
    deck = Presentation(exports["pptx"])
    assert len(deck.slides) == len(book.pages)
    assert book.pages[2].source_text in deck.slides[2].notes_slide.notes_text_frame.text
    assert "https://www.youtube.com/watch?v=Tcn5Yb2K0h4&t=72" in deck.slides[2].notes_slide.notes_text_frame.text
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in deck.slides[2].shapes)
    links = [run.hyperlink.address for shape in deck.slides[2].shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.hyperlink.address]
    assert links == ["https://www.youtube.com/watch?v=Tcn5Yb2K0h4&t=72"]
    pdf = PdfReader(exports["pdf"])
    assert len(pdf.pages) == len(book.pages)
    assert "/XObject" in pdf.pages[2]["/Resources"]
    assert "Transcript-only page" in pdf.pages[3].extract_text()
    html = exports["html"].read_text()
    assert html.count('<article class="book-page"') == len(book.pages)
    assert "data:image/png;base64," in html
    assert "Captured source visual for interval 1:12–2:00" in html
    assert "READING TARGET" in html
    assert "AT A GLANCE" in html
    assert "PRESENTER SUMMARY" in html
    assert "Reading target: 60 minutes" in html
    assert "About 60 minutes to read" not in html
    assert "&lt;script&gt;alert(&#x27;never execute&#x27;)&lt;/script&gt;" in html
    assert "<script>alert('never execute')</script>" not in html
    assert "complete source transcript " * 300 in html
    assert "&lt;with&gt; examples &amp; sources" in html
    manifest = json.loads(exports["json"].read_text())
    assert len(manifest["pages"]) == len(book.pages)
    assert manifest["pages"][2]["source_text"] == book.pages[2].source_text
    assert any("transcript-only" in warning for warning in manifest["warnings"])


def test_extractive_label_and_missing_image_warnings_are_honest_and_compact(tmp_path):
    pages = [Page(i + 1, f"Excerpt {i}", i * 20, i * 20 + 20, "The exact source wording.", [], "", source_text="The exact source wording.", uncertainty="No AI summary was generated.") for i in range(198)]
    pages.append(Page(199, "Missing captured file", 4000, 4020, "A summary", [], "", frame_path=str(tmp_path / "missing.png")))
    book = Book("Transcript-only book", "https://youtu.be/abc", 4020, pages)
    output = export_book(book, tmp_path / "extractive", ("html",))
    html = output["html"].read_text()
    assert html.count("TRANSCRIPT EXCERPT") == 198
    assert "198 content page(s)" in book.warnings[1]
    assert len(book.warnings) == 2
    assert "Page 199: captured image file is missing" in book.warnings[0]


def test_html_only_never_links_unsafe_scheme(tmp_path):
    book = sample_book(tmp_path)
    book.source_url = 'javascript:alert("bad")'
    exports = export_book(book, tmp_path / "html", ("html",))
    html = exports["html"].read_text()
    assert 'href="javascript:' not in html
    assert set(exports) == {"html", "json"}


def test_long_text_is_explicitly_shortened_and_full_text_retained(tmp_path):
    from pptx import Presentation
    from pypdf import PdfReader

    title = "A very long title with many ideas and details " * 100
    summary = "A presenter explanation that cannot fit in a narrow sidebar. " * 300
    book = Book("Long title test", "https://youtu.be/abc", 100, [Page(1, title, 0, 100, summary, ["A long key idea " * 200], "The important conclusion " * 200, source_text="Full underlying transcript")])
    exports = export_book(book, tmp_path / "long")
    assert any("shortened" in warning for warning in book.warnings)
    slide = Presentation(exports["pptx"]).slides[0]
    assert summary in slide.notes_slide.notes_text_frame.text
    slide_text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert "Full text is available in HTML and PowerPoint notes." in slide_text
    assert "Full text is available in HTML and PowerPoint notes." in PdfReader(exports["pdf"]).pages[0].extract_text()
    assert summary in exports["html"].read_text()
    assert title in exports["html"].read_text()


def test_rejects_unknown_format_empty_book_and_invalid_interval_before_writing(tmp_path):
    book = sample_book(tmp_path)
    with pytest.raises(ValueError, match="Unsupported export"):
        export_book(book, tmp_path / "bad", ("docx",))
    assert not (tmp_path / "bad").exists()
    book.pages = []
    with pytest.raises(ValueError, match="at least one page"):
        export_book(book, tmp_path / "bad")
    book.pages = [Page(1, "Bad range", 30, 20, "", [], "")]
    with pytest.raises(ValueError, match="invalid source interval"):
        export_book(book, tmp_path / "bad")
